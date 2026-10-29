"""Onboarding for the dashboard (tab "Început"): from a CV to MASTER_SKILLS_POOL.md and ANSWERS_POOL.json.

Steps, each one a file in <workspace>/onboarding/ (the state is read from the files, so a restart loses nothing):
    1 cv          upload a PDF / DOCX (+ optional LinkedIn / free text) → cv.<ext>, cv.txt, extra.txt
                  (no pypdf installed: no cv.txt, step 2 has Claude read cv.pdf itself)
    2 facts       Claude task `onboard_extract` reads them (setup-profile rules) → facts.json (proposed, not confirmed)
    3 confirmed   the user corrects the fixed facts and the skill levels on the page → facts.json "confirmed": true
    4 answers     the form-answer form (contact, screening, work mode, consents, ...) → ANSWERS_POOL.json + answers.json
    5 enrichment  the guided chat (2 rounds of at most 6 questions: metrics, production status, interview checks)
                  → enrichment.md (Claude writes it); may be skipped
    6 written     Claude task `onboard_write` → MASTER_SKILLS_POOL.md, profile.json, AGENTS.md, the rest of
                  ANSWERS_POOL.json, scoring.json keys (snapshot: revertable)
Nothing here calls Claude; the Claude steps are runner tasks and the chat is dashboard.mirel.OnboardingGuide.
"""
import base64
import datetime as dt
import io
import json
import os
import re

from jobpilot.score import DEFAULT_DOMAIN   # also the example scoring.json's placeholder

from . import profile

MAX_BYTES = 10 * 1024 * 1024
MIN_TEXT = 200        # less than this from a PDF: a scan without a text layer
FACT_LISTS = ("roles", "education", "certifications", "languages", "skills", "projects")
LEVELS = ("A", "B", "C", "D", "E")
CONTACT_KEYS = ("first_name", "last_name", "full_name", "email", "phone", "phone_country_code", "city", "country",
                "linkedin", "github", "portfolio")
SCREENING_KEYS = ("current_company", "current_title", "work_authorization", "requires_visa_sponsorship",
                  "how_did_you_hear")
VOLUNTARY_KEYS = ("gender", "ethnicity", "veteran_status", "disability")
CONSENTS = ("this_application_data", "future_recruitment", "marketing")
CV_LANGUAGES = ("EN", "RO", "DE", "FR", "ES", "IT")


def folder(ws):
    return ws.rel("onboarding")


def _path(ws, name):
    return os.path.join(folder(ws), name)


def _read_json(path, default=None):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as f:
        return json.load(f)


def _write_json(path, data):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as f:
        json.dump(data, f, ensure_ascii=False, indent=2)
        f.write("\n")
    os.replace(tmp, path)


def state(ws):
    """→ {steps: {cv, facts, confirmed, answers, enrichment, written}, pool_exists, needs_onboarding, cv: {name,
    chars}, facts, answers}"""
    facts = _read_json(_path(ws, "facts.json"))
    answers = _read_json(_path(ws, "answers.json"))
    cv_text = _path(ws, "cv.txt")
    cv_name = next((f for f in sorted(os.listdir(folder(ws))) if re.match(r"cv\.(pdf|docx)$", f)), None) \
        if os.path.isdir(folder(ws)) else None
    pool = os.path.exists(ws.path("skills_pool"))
    written = os.path.exists(_path(ws, "written.json"))
    return {
        "steps": {"cv": cv_name is not None, "facts": facts is not None,
                  "confirmed": bool(facts and facts.get("confirmed")), "answers": answers is not None,
                  "enrichment": os.path.exists(_path(ws, "enrichment.md")), "written": written},
        "pool_exists": pool, "needs_onboarding": not pool,
        "cv": {"name": cv_name, "chars": len(open(cv_text, encoding="utf-8").read()) if os.path.exists(cv_text) else 0,
               "text": os.path.exists(cv_text)},   # False: Claude reads the PDF itself at step 2
        "facts": facts, "answers": answers or defaults(ws, facts),
        "levels": LEVELS, "cv_languages": CV_LANGUAGES,
    }


def extract_text(filename, data):
    """PDF / DOCX bytes → plain text (pypdf / python-docx); None for a PDF without pypdf installed. ValueError for
    other files or a DOCX without python-docx."""
    ext = os.path.splitext(filename.lower())[1]
    if ext == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError:
            return None   # Claude reads the PDF at step 2
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages).strip()
    if ext == ".docx":
        try:
            import docx
        except ImportError:
            raise ValueError('python-docx is not installed: pip install "./cv-assistant[cv]"')
        doc = docx.Document(io.BytesIO(data))
        lines = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                lines.append(" | ".join(c.text.strip() for c in row.cells))
        return "\n".join(l for l in lines if l.strip()).strip()
    raise ValueError("the CV must be a PDF or a DOCX file")


def upload_cv(ws, filename, content_b64, extra_text=""):
    """Saves the CV and its text (and the optional extra text) → {name, chars}. ValueError on a bad file."""
    filename = os.path.basename(filename or "")
    ext = os.path.splitext(filename.lower())[1]
    if ext not in (".pdf", ".docx"):
        raise ValueError("the CV must be a PDF or a DOCX file")
    try:
        data = base64.b64decode(content_b64 or "", validate=True)
    except (ValueError, TypeError):
        raise ValueError("the file did not arrive whole; choose it again")
    if not data:
        raise ValueError("the file is empty")
    if len(data) > MAX_BYTES:
        raise ValueError("the CV is larger than 10 MB")
    text = extract_text(filename, data)
    if text is not None and len(text) < MIN_TEXT:
        raise ValueError("almost no text could be read from this file (a scanned PDF?); upload the DOCX or a PDF "
                         "exported from Word")
    os.makedirs(folder(ws), exist_ok=True)
    for old in ("cv.pdf", "cv.docx"):
        if os.path.exists(_path(ws, old)):
            os.remove(_path(ws, old))
    with open(_path(ws, "cv" + ext), "wb") as f:
        f.write(data)
    if text is None:
        if os.path.exists(_path(ws, "cv.txt")):
            os.remove(_path(ws, "cv.txt"))
    else:
        with open(_path(ws, "cv.txt"), "w", encoding="utf-8") as f:
            f.write(text + "\n")
    with open(_path(ws, "extra.txt"), "w", encoding="utf-8") as f:
        f.write((extra_text or "").strip()[:30000] + "\n")
    for stale in ("facts.json", "written.json"):   # a new CV is read again from the start
        if os.path.exists(_path(ws, stale)):
            os.remove(_path(ws, stale))
    return {"name": "cv" + ext, "chars": len(text or "")}


def _clean(value, limit=500):
    return re.sub(r"\s+", " ", str(value or "")).strip()[:limit]


def save_facts(ws, facts):
    """The user's corrected facts → facts.json, confirmed. ValueError when the shape is wrong."""
    if not isinstance(facts, dict):
        raise ValueError("facts must be an object")
    current = _read_json(_path(ws, "facts.json"))
    if current is None:
        raise ValueError("read the CV first (step 2)")
    out = dict(current)
    for key in ("name", "headline"):
        if key in facts:
            out[key] = _clean(facts[key], 200)
    if "contact" in facts:
        if not isinstance(facts["contact"], dict):
            raise ValueError("contact must be an object")
        out["contact"] = {k: _clean(v, 300) for k, v in facts["contact"].items()}
    for key in FACT_LISTS:
        if key in facts:
            items = facts[key]
            if not isinstance(items, list) or not all(isinstance(i, dict) for i in items):
                raise ValueError(f"{key} must be a list of objects")
            out[key] = [{k: (_clean(v, 1000) if not isinstance(v, list) else [_clean(x, 500) for x in v])
                         for k, v in i.items()} for i in items]
    for s in out.get("skills") or []:
        level = (s.get("level") or "").strip().upper()[:1]
        if level not in LEVELS:
            raise ValueError(f"skill {s.get('name')!r}: the level must be one of {', '.join(LEVELS)}")
        s["level"] = level
    out["confirmed"] = True
    out["confirmed_at"] = dt.date.today().isoformat()
    _write_json(_path(ws, "facts.json"), out)
    return state(ws)


def defaults(ws, facts=None):
    """The form-answer form's starting values: what the answers pool already has, else the facts from the CV."""
    pool = profile._load(ws.path("answers_pool"))
    contact = dict(pool.get("contact") or {})
    fc = (facts or {}).get("contact") or {}
    name = (facts or {}).get("name") or ""
    guess = {"full_name": name, "first_name": name.split(" ")[0] if name else "",
             "last_name": " ".join(name.split(" ")[1:]) if " " in name else "", "email": fc.get("email"),
             "phone": fc.get("phone"), "city": fc.get("city") or fc.get("location"), "country": fc.get("country"),
             "linkedin": fc.get("linkedin"), "github": fc.get("github"), "portfolio": fc.get("portfolio")}
    for k in CONTACT_KEYS:
        contact[k] = contact.get(k) or guess.get(k) or ""
    screening = pool.get("screening") or {}
    roles = (facts or {}).get("roles") or []
    current = roles[0] if roles else {}   # newest first (facts.json schema)
    prof = profile.get(ws)
    consents = {c["id"]: c.get("answer") for c in pool.get("consents") or [] if c.get("id") in CONSENTS}
    return {
        "contact": contact,
        "screening": {
            **{k: screening.get(k, "") for k in SCREENING_KEYS},
            "current_company": screening.get("current_company") or (current or {}).get("employer", ""),
            "current_title": screening.get("current_title") or (current or {}).get("title", ""),
            "how_did_you_hear": screening.get("how_did_you_hear") or "Company website / job board",
        },
        "profile": {**{k: prof[k] for k in ("workplaces", "preferred_contract_type", "notice_period",
                                            "earliest_start_date", "willing_to_relocate", "salary")},
                    # none set yet: the roles Claude guessed from the CV, for the user to keep or change
                    "target_titles": prof["target_titles"] or
                    [t for t in (facts or {}).get("target_roles_guess") or [] if t][:profile.MAX_TITLES]},
        "voluntary": {k: (pool.get("voluntary_disclosures") or {}).get(k) or "Prefer not to say" for k in VOLUNTARY_KEYS},
        "consents": {k: consents.get(k) for k in CONSENTS},
        "cv_languages": (facts or {}).get("cv_languages") or ["EN"],
        # the example config's placeholder is no answer
        "target_domain": "" if prof.get("target_domain", "") in ("", DEFAULT_DOMAIN) else prof["target_domain"],
    }


def save_answers(ws, answers):
    """The form-answer form → ANSWERS_POOL.json (contact, screening, voluntary disclosures, consents) + the work
    settings via profile.save + onboarding/answers.json. ValueError on bad input."""
    if not isinstance(answers, dict):
        raise ValueError("answers must be an object")
    contact = {k: _clean(v, 300) for k, v in (answers.get("contact") or {}).items() if k in CONTACT_KEYS}
    if not contact.get("email") or "@" not in contact["email"]:
        raise ValueError("an email address is needed (it goes into every application form)")
    if not contact.get("full_name"):
        contact["full_name"] = " ".join(x for x in (contact.get("first_name"), contact.get("last_name")) if x)
    languages = [l for l in answers.get("cv_languages") or [] if l in CV_LANGUAGES]
    if not languages:
        raise ValueError("choose at least one CV language")
    consents = answers.get("consents") or {}
    for k in CONSENTS:
        if consents.get(k) not in (True, False, None):
            raise ValueError(f"consent {k}: yes, no or not set")
    work = answers.get("profile") or {}
    profile.save(ws, {k: v for k, v in work.items() if k in ("workplaces", "preferred_contract_type", "notice_period",
                                                             "earliest_start_date", "willing_to_relocate",
                                                             "target_titles", "salary")})
    path = ws.path("answers_pool")
    pool = profile._load(path)
    pool.setdefault("contact", {}).update(contact)
    screening = pool.setdefault("screening", {})
    for k, v in (answers.get("screening") or {}).items():
        if k in SCREENING_KEYS:
            screening[k] = _clean(v, 300)
    voluntary = pool.setdefault("voluntary_disclosures", {})
    for k, v in (answers.get("voluntary") or {}).items():
        if k in VOLUNTARY_KEYS:
            voluntary[k] = _clean(v, 100) or "Prefer not to say"
    for c in pool.get("consents") or []:
        if c.get("id") in CONSENTS and c["id"] in consents:
            c["answer"] = consents[c["id"]]   # only what the user ticked: true / false, None = ask in the form
    profile._save(path, pool)
    if answers.get("target_domain") is not None:
        scoring = profile._load(ws.path("scoring"))
        scoring["target_domain"] = _clean(answers["target_domain"], 200)
        profile._save(ws.path("scoring"), scoring)
    saved = {**answers, "contact": contact, "cv_languages": languages}
    _write_json(_path(ws, "answers.json"), saved)
    return state(ws)


def skip_enrichment(ws):
    """The user skips the guided chat: an enrichment.md that says so (step 6 then uses the CV facts only)."""
    path = _path(ws, "enrichment.md")
    if not os.path.exists(path):
        os.makedirs(folder(ws), exist_ok=True)
        with open(path, "w", encoding="utf-8") as f:
            f.write("# Enrichment\n\nSkipped by the user: no extra answers; unknown metrics go to "
                    "\"Information to collect\".\n")
    return state(ws)


def mark_written(ws, task_id):
    _write_json(_path(ws, "written.json"), {"task_id": task_id, "at": dt.datetime.now().isoformat(timespec="seconds")})


def facts_summary(ws):
    """Short text for the guide chat's context: the confirmed facts and the open questions."""
    facts = _read_json(_path(ws, "facts.json")) or {}
    answers = _read_json(_path(ws, "answers.json")) or {}
    return json.dumps({"facts": {k: facts.get(k) for k in ("name", "headline", "roles", "skills", "projects",
                                                           "open_questions")},
                       "target_titles": (answers.get("profile") or {}).get("target_titles")},
                      ensure_ascii=False, indent=1)
